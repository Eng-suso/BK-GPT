"""Acquisizione strutturale di un workbook XLSX: l'Evidence Bucket per Excel.

Il parser e' openpyxl (MIT), il lettore di riferimento per il formato: qui
non si interpreta niente, si traduce il workbook in evidenze ancorate. Ogni
cella con un contenuto diventa un'evidenza `Foglio!B7` con il suo tipo, il
valore salvato, la formula che lo produce, il formato numerico, l'unione in
cui sta, la tabella e l'intestazione di colonna a cui appartiene, e se e'
nascosta. I commenti di cella sono evidenze a parte.

Cosa openpyxl non legge (caselle di testo nei disegni) o legge solo come
valore salvato (formule che puntano a file esterni) diventa un
`AcquisitionIssue`: la fonte dichiara cosa manca, non lo perde in silenzio.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import warnings
import zipfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from backend.workspace_services.evidence.canonical import (
    AcquisitionIssue,
    Anchor,
    CanonicalSource,
    EvidenceSegment,
)

PARSER_ID = "openpyxl-structural/1"

MAX_ARCHIVE_BYTES = 40 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 5_000
MAX_CELLS = 1_000_000
_CFB_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_PLAIN_SHEET_NAME = re.compile(r"[A-Za-z_À-￿][\w.À-￿]*")
_LOOKS_LIKE_CELL = re.compile(r"[A-Za-z]{1,3}\d+")


class WorkbookUnreadable(ValueError):
    """Il file non e' un workbook che si puo' aprire."""


def sheet_ref(sheet: str, coordinate: str) -> str:
    """Il riferimento come lo scrive Excel: `Ordini!B7`, `'Ciclo passivo'!B7`."""
    if _PLAIN_SHEET_NAME.fullmatch(sheet) and not _LOOKS_LIKE_CELL.fullmatch(sheet):
        return f"{sheet}!{coordinate}"
    return "'" + sheet.replace("'", "''") + f"'!{coordinate}"


def parse_xlsx(payload: bytes) -> CanonicalSource:
    """Traduce un workbook nelle sue evidenze ancorate.

    Args:
        payload: I byte del file, non fidati.

    Returns:
        La rappresentazione canonica, con `status == "partial"` quando una
        parte del workbook non e' stata acquisita.

    Raises:
        WorkbookUnreadable: Il file e' cifrato, non e' un archivio OOXML
            valido, supera i limiti o contiene XML non sicuro.
    """
    issues: list[AcquisitionIssue] = []
    check_ooxml_archive(payload, "Excel")

    try:
        with warnings.catch_warnings():
            # openpyxl avvisa per estensioni che non legge (validazione dati,
            # formattazione condizionale): non sono contenuto della fonte.
            warnings.simplefilter("ignore")
            formulas = load_workbook(io.BytesIO(payload), data_only=False)
            values = load_workbook(io.BytesIO(payload), data_only=True)
    except WorkbookUnreadable:
        raise
    except Exception as exc:  # noqa: BLE001 - qualunque rottura del file e' "non leggibile"
        raise WorkbookUnreadable("Il file Excel non può essere letto.") from exc

    if getattr(formulas, "_external_links", None):
        issues.append(
            AcquisitionIssue(
                code="external_links",
                severity="info",
                message="Il workbook è collegato ad altri file: i valori di quelle formule sono quelli salvati.",
            )
        )

    segments: list[EvidenceSegment] = []
    sheets: list[dict[str, Any]] = []
    cell_count = 0
    truncated = False

    for index, sheet in enumerate(formulas.worksheets):
        if truncated:
            break
        cached = values[sheet.title]
        hidden_sheet = sheet.sheet_state != "visible"
        hidden_rows = sorted(row for row, dim in sheet.row_dimensions.items() if dim.hidden)
        hidden_columns = _hidden_columns(sheet)
        merged = [str(item) for item in sheet.merged_cells.ranges]
        merged_top_left = {
            (item.min_row, item.min_col): str(item) for item in sheet.merged_cells.ranges
        }
        tables = _tables(sheet)
        sheets.append(
            {
                "name": sheet.title,
                "index": index,
                "state": sheet.sheet_state,
                "dimension": sheet.dimensions,
                "hidden_rows": hidden_rows,
                "hidden_columns": [get_column_letter(item) for item in sorted(hidden_columns)],
                "merged": merged,
                "tables": tables,
            }
        )

        # `_cells` e' l'insieme sparso delle celle esistenti. `iter_rows`
        # creerebbe la griglia densa fino all'ultima cella: un solo valore
        # perso in XFD1048576 basta a farla esplodere.
        for row, column in sorted(sheet._cells):
            cell = sheet._cells[(row, column)]
            raw = cell.value
            if raw is None and (row, column) not in merged_top_left:
                continue
            cell_count += 1
            if cell_count > MAX_CELLS:
                issues.append(
                    AcquisitionIssue(
                        code="cell_limit",
                        severity="partial",
                        message=f"Il workbook supera {MAX_CELLS:,} celle: le successive non sono state acquisite.".replace(",", "."),
                        ref=sheet_ref(sheet.title, cell.coordinate),
                    )
                )
                truncated = True
                break

            segment = _cell_segment(
                sheet=sheet.title,
                sheet_index=index,
                cell=cell,
                cached=cached.cell(row=row, column=column).value,
                issues=issues,
            )
            if segment is None:
                continue
            attributes = segment.attributes
            if hidden_sheet:
                attributes["sheet_hidden"] = True
            if row in hidden_rows or column in hidden_columns:
                attributes["hidden"] = True
            if (row, column) in merged_top_left:
                attributes["merged_range"] = merged_top_left[(row, column)]
            for table in tables:
                role = _table_role(table, row, column)
                if role is None:
                    continue
                attributes["table"] = table["name"]
                attributes["table_role"] = role[0]
                if role[1] is not None:
                    attributes["table_column"] = role[1]
                break
            segments.append(segment)

            if cell.comment is not None and (cell.comment.text or "").strip():
                segments.append(
                    EvidenceSegment(
                        anchor=Anchor(
                            kind="cell_comment",
                            ref=sheet_ref(sheet.title, cell.coordinate) + "#commento",
                            locator={"sheet": sheet.title, "sheet_index": index, "row": row, "column": column, "cell": cell.coordinate},
                        ),
                        text=cell.comment.text.strip(),
                        value_type="text",
                        value=cell.comment.text.strip(),
                        attributes={"author": cell.comment.author} if cell.comment.author else {},
                    )
                )

    issues.extend(_unread_drawing_text(payload))

    return CanonicalSource(
        format="xlsx",
        parser=PARSER_ID,
        segments=tuple(segments),
        structure={
            "sheets": sheets,
            "defined_names": _defined_names(formulas),
            "date_system": 1904 if formulas.epoch.year == 1904 else 1900,
        },
        issues=tuple(issues),
    )


def render_text(source: CanonicalSource) -> str:
    """Il testo leggibile del workbook, con le ancore a vista.

    Una riga per riga di foglio, ogni valore preceduto dalla sua colonna:
    `7: A=Ordine 42 | C=1200 (=SUM(C2:C6))`. Una colonna vuota non fa
    scivolare le altre, perche' ogni valore porta la sua lettera.
    """
    blocks: list[str] = []
    for sheet in source.structure.get("sheets", []):
        name = sheet["name"]
        header = f"[{name}]" + (" (foglio nascosto)" if sheet["state"] != "visible" else "")
        rows: dict[int, list[str]] = {}
        comments: list[str] = []
        for segment in source.segments:
            locator = segment.anchor.locator
            if locator.get("sheet") != name:
                continue
            letter = get_column_letter(locator["column"])
            if segment.anchor.kind == "cell_comment":
                author = segment.attributes.get("author")
                comments.append(
                    f"Commento {locator['cell']}" + (f" ({author})" if author else "") + f": {segment.text}"
                )
                continue
            text = segment.text
            formula = segment.attributes.get("formula")
            if formula:
                text = f"{text} (={formula})" if text else f"(={formula})"
            rows.setdefault(locator["row"], []).append(f"{letter}={text}")
        lines = [f"{row}: " + " | ".join(cells) for row, cells in sorted(rows.items())]
        blocks.append("\n".join([header, *lines, *comments]))
    return "\n\n".join(blocks)


def check_ooxml_archive(payload: bytes, label: str) -> None:
    """Rifiuta un file Office cifrato, sproporzionato o con XML non sicuro.

    Vale per ogni formato OOXML (xlsx, docx, pptx): sono archivi zip di XML,
    e un DOCTYPE con entita' e' la porta di un XXE o di una bomba di espansione.
    """
    if payload.startswith(_CFB_MAGIC):
        raise WorkbookUnreadable(
            f"Il file {label} è protetto da password o in un formato vecchio: salvalo senza password nel formato attuale."
        )
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as exc:
        raise WorkbookUnreadable(f"Il file {label} non è valido.") from exc
    entries = archive.infolist()
    if len(entries) > MAX_ARCHIVE_ENTRIES:
        raise WorkbookUnreadable(f"Il file {label} contiene troppi elementi.")
    if sum(item.file_size for item in entries) > MAX_ARCHIVE_BYTES:
        raise WorkbookUnreadable(f"Il contenuto del file {label} è troppo grande.")
    for item in entries:
        if not item.filename.endswith((".xml", ".rels")):
            continue
        head = archive.read(item)[:4096].upper()
        if b"<!DOCTYPE" in head or b"<!ENTITY" in head:
            raise WorkbookUnreadable(f"Il file {label} contiene XML non sicuro.")


def _cell_segment(
    *,
    sheet: str,
    sheet_index: int,
    cell: Any,
    cached: Any,
    issues: list[AcquisitionIssue],
) -> EvidenceSegment | None:
    raw = cell.value
    ref = sheet_ref(sheet, cell.coordinate)
    attributes: dict[str, Any] = {}
    if cell.number_format and cell.number_format != "General":
        attributes["number_format"] = cell.number_format

    value = raw
    if isinstance(raw, ArrayFormula):
        attributes["formula"] = (raw.text or "").removeprefix("=")
        attributes["formula_kind"] = "array"
        attributes["array_range"] = raw.ref
        value = cached
    elif isinstance(raw, DataTableFormula):
        attributes["formula_kind"] = "data_table"
        value = cached
    elif cell.data_type == "f" and isinstance(raw, str):
        attributes["formula"] = raw.removeprefix("=")
        attributes["formula_kind"] = "normal"
        value = cached

    if "formula_kind" in attributes and value is None:
        issues.append(
            AcquisitionIssue(
                code="formula_not_calculated",
                severity="info",
                message="La formula non ha un valore salvato: il file non è stato ricalcolato prima del salvataggio.",
                ref=ref,
            )
        )

    value_type, typed, text = _typed(value, cell.data_type if "formula_kind" not in attributes else None)
    if value_type == "empty" and "formula_kind" not in attributes:
        # Solo la cella in alto a sinistra di un'unione vuota arriva qui:
        # l'unione resta nella struttura del foglio, non e' un'evidenza.
        return None

    return EvidenceSegment(
        anchor=Anchor(
            kind="cell",
            ref=ref,
            locator={
                "sheet": sheet,
                "sheet_index": sheet_index,
                "row": cell.row,
                "column": cell.column,
                "cell": cell.coordinate,
            },
        ),
        text=text,
        value_type=value_type,
        value=typed,
        attributes=attributes,
    )


def _typed(value: Any, data_type: str | None) -> tuple[str, Any, str]:
    """Tipo, valore serializzabile e testo mostrato, per un valore di cella."""
    if value is None:
        return "empty", None, ""
    if isinstance(value, bool):
        return "boolean", value, "VERO" if value else "FALSO"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return "date", value.date().isoformat(), value.date().isoformat()
        return "datetime", value.isoformat(), value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return "date", value.isoformat(), value.isoformat()
    if isinstance(value, dt.time):
        return "time", value.isoformat(), value.isoformat()
    if isinstance(value, dt.timedelta):
        seconds = value.total_seconds()
        hours, remainder = divmod(int(round(seconds)), 3600)
        minutes, secs = divmod(remainder, 60)
        return "duration", seconds, f"{hours}:{minutes:02d}:{secs:02d}"
    if isinstance(value, int):
        return "number", value, str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return "number", int(value), str(int(value))
        return "number", value, repr(value)
    text = str(value)
    # Una cella scritta a mano "#N/A" e' testo; e' errore solo se Excel lo
    # dice (`data_type == "e"`) o se e' il risultato salvato di una formula,
    # dove openpyxl non porta il tipo (`data_type is None`).
    if data_type == "e" or (data_type is None and text.rstrip("!?0/").upper() in _ERRORS):
        return "error", text, text
    return "text", text, text


_ERRORS = {"#NULL", "#DIV", "#VALUE", "#REF", "#NAME", "#NUM", "#N/A", "#GETTING_DATA", "#SPILL", "#CALC"}


def _hidden_columns(sheet: Any) -> set[int]:
    hidden: set[int] = set()
    for dimension in sheet.column_dimensions.values():
        if not dimension.hidden:
            continue
        # Una dimensione copre un intervallo di colonne (min..max).
        low = dimension.min or 0
        high = dimension.max or low
        hidden.update(range(low, high + 1))
    return hidden


def _tables(sheet: Any) -> list[dict[str, Any]]:
    tables: list[dict[str, Any]] = []
    for table in sheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        tables.append(
            {
                "name": table.displayName or table.name,
                "ref": table.ref,
                "header_rows": table.headerRowCount if table.headerRowCount is not None else 1,
                "totals_rows": table.totalsRowCount or 0,
                "columns": [column.name for column in table.tableColumns],
                "bounds": [min_col, min_row, max_col, max_row],
            }
        )
    return tables


def _table_role(table: dict[str, Any], row: int, column: int) -> tuple[str, str | None] | None:
    min_col, min_row, max_col, max_row = table["bounds"]
    if not (min_col <= column <= max_col and min_row <= row <= max_row):
        return None
    name = table["columns"][column - min_col] if column - min_col < len(table["columns"]) else None
    if row < min_row + table["header_rows"]:
        return "header", name
    if row > max_row - table["totals_rows"]:
        return "total", name
    return "data", name


def _defined_names(workbook: Any) -> dict[str, str]:
    names: dict[str, str] = {}
    try:
        for name, definition in workbook.defined_names.items():
            names[name] = definition.attr_text
    except Exception:  # noqa: BLE001 - un nome definito malformato non e' contenuto
        pass
    return names


def _unread_drawing_text(payload: bytes) -> list[AcquisitionIssue]:
    """openpyxl scarta le forme dei disegni: se contengono testo, la fonte e' parziale."""
    issues: list[AcquisitionIssue] = []
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        for item in archive.infolist():
            if not re.match(r"^xl/drawings/drawing\d+\.xml$", item.filename):
                continue
            content = archive.read(item)
            if b":t>" in content or b"<t>" in content:
                issues.append(
                    AcquisitionIssue(
                        code="drawing_text_not_read",
                        severity="partial",
                        message="Il workbook contiene caselle di testo o forme con testo che non sono state acquisite.",
                        ref=item.filename,
                    )
                )
    return issues
