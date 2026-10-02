from __future__ import annotations

import uuid
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient

from backend.app import app
from backend.settings import settings


pytestmark = pytest.mark.skipif(
    not settings.workspace_database_url, reason="serve WORKSPACE_DATABASE_URL"
)


@pytest.fixture()
def tenant() -> str:
    """Un tenant per test: la coda delle fonti e' condivisa, il worker no."""
    return f"ingest-{uuid.uuid4().hex[:10]}"


@pytest.fixture()
def http(tenant: str):
    with TestClient(app, headers={"X-DeliR-Tenant-ID": tenant}) as client:
        yield client


def _drain(source_worker, tenant: str) -> None:
    while source_worker.drain_once(10, only_tenant_id=tenant):
        pass


def _project(http: TestClient) -> tuple[dict, dict]:
    client = http.post(
        "/v1/workspace/clients", json={"name": f"Ingestion {uuid.uuid4().hex[:8]}"}
    ).json()
    project = http.post(
        "/v1/workspace/projects",
        json={"client_id": client["id"], "name": "Acquisti"},
    ).json()
    process = http.post(
        f"/v1/workspace/projects/{project['id']}/processes",
        json={"name": "Procure to pay"},
    ).json()
    return project, process


def test_upload_preserves_dimensions_and_makes_text_readable(http: TestClient, tenant: str):
    project, process = _project(http)

    response = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={
            "roles": '["process_evidence", "policy"]',
            "retention": "persistent",
            "scopes": f'[{{"type":"project","id":"{project["id"]}"}},'
            f'{{"type":"process","id":"{process["id"]}"}}]',
        },
        files={
            "file": (
                "Procedura_Acquisti.md",
                b"# Acquisti\nIl CFO approva gli ordini sopra EUR 30.000.",
                "text/markdown",
            )
        },
    )

    assert response.status_code == 201, response.text
    source = response.json()
    # Il caricamento non legge il file: lo conserva e lo mette in coda.
    assert source["acquisition_status"] == "pending"
    assert source["name"] == "Procedura_Acquisti.md"
    assert source["roles"] == ["process_evidence", "policy"]
    assert source["retention"] == "persistent"
    assert source["status"] == "extracted"
    assert source["byte_size"] > 0
    assert source["content_hash"]
    assert source["scopes"] == [
        {"type": "project", "id": project["id"]},
        {"type": "process", "id": process["id"]},
    ]

    from backend.workers import source_worker

    # Un file non ancora letto non puo' diventare evidenza.
    early = http.post(f"/v1/workspace/sources/{source['id']}/verify")
    assert early.status_code == 409

    _drain(source_worker, tenant)
    listed = {item["id"]: item for item in http.get(f"/v1/workspace/projects/{project['id']}/sources").json()}
    assert listed[source["id"]]["acquisition_status"] == "done"

    evidence = http.get(f"/v1/workspace/sources/{source['id']}/evidence")
    assert evidence.status_code == 200
    assert evidence.json() == []  # Markdown non ha ancora un parser strutturale

    document = http.get(f"/v1/workspace/sources/{source['id']}/document")
    assert document.status_code == 200
    assert "Il CFO approva" in document.json()["content"]
    assert document.json()["has_content"] is True

    original = http.get(f"/v1/workspace/sources/{source['id']}/original")
    assert original.status_code == 200
    assert original.content.startswith(b"# Acquisti")
    assert http.get(
        f"/v1/workspace/sources/{source['id']}/original",
        headers={"X-DeliR-Tenant-ID": "tenant-altro"},
    ).status_code == 404

    from backend.graphs.process.nodes import load_evidence_ledger
    from backend.security import reset_current_tenant_id, set_current_tenant_id

    def ledger_ids() -> list[str]:
        token = set_current_tenant_id(tenant)
        try:
            return load_evidence_ledger(project["id"], process["id"])["source_ids"]
        finally:
            reset_current_tenant_id(token)

    # Letto ma non confermato: non e' ancora evidenza del processo.
    assert source["id"] not in ledger_ids()
    verified = http.post(f"/v1/workspace/sources/{source['id']}/verify")
    assert verified.status_code == 200
    assert verified.json()["status"] == "approved"
    assert source["id"] in ledger_ids()

    deleted = http.delete(f"/v1/workspace/projects/{project['id']}")
    assert deleted.status_code == 200
    assert http.get(f"/v1/workspace/sources/{source['id']}/original").status_code == 404


def test_upload_rejects_unsupported_files_and_foreign_process_scope(http: TestClient):
    project, _ = _project(http)
    other_project, other_process = _project(http)

    unsupported = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={"roles": '["context"]', "retention": "persistent", "scopes": "[]"},
        files={"file": ("payload.exe", b"MZ", "application/octet-stream")},
    )
    assert unsupported.status_code == 415

    temporary = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={"roles": '["context"]', "retention": "temporary", "scopes": "[]"},
        files={"file": ("nota.txt", b"testo", "text/plain")},
    )
    assert temporary.status_code == 400
    assert "flusso di contesto" in str(temporary.json())

    foreign_scope = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={
            "roles": '["process_evidence"]',
            "retention": "persistent",
            "scopes": f'[{{"type":"process","id":"{other_process["id"]}"}}]',
        },
        files={"file": ("nota.txt", b"testo", "text/plain")},
    )
    assert foreign_scope.status_code == 400
    assert other_project["id"] != project["id"]

    malformed_roles = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={"roles": '[{"unexpected":true}]', "retention": "persistent", "scopes": "[]"},
        files={"file": ("nota.txt", b"testo", "text/plain")},
    )
    assert malformed_roles.status_code == 400


def test_upload_is_idempotent_for_the_same_file_and_dimensions(http: TestClient, tenant: str):
    project, _ = _project(http)
    fields = {
        "roles": '["context"]',
        "retention": "persistent",
        "scopes": f'[{{"type":"project","id":"{project["id"]}"}}]',
    }

    first = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data=fields,
        files={"file": ("contesto.txt", b"stesso contenuto", "text/plain")},
    )
    second = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data=fields,
        files={"file": ("contesto.txt", b"stesso contenuto", "text/plain")},
    )

    assert first.status_code == 201
    assert second.status_code == 200
    assert second.json()["id"] == first.json()["id"]

    from backend.security import reset_current_tenant_id, set_current_tenant_id
    from backend.workspace_database import create_ingested_source

    def create_once() -> str:
        token = set_current_tenant_id(tenant)
        try:
            source, _ = create_ingested_source(
                project_id=project["id"],
                name="concorrenza.txt",
                roles=["context"],
                retention="persistent",
                scopes=[{"type": "project", "id": project["id"]}],
                content_hash="a" * 64,
                byte_size=5,
                mime_type="text/plain",
                storage_key="source_uploads/test/concorrenza.txt",
            )
            return source["id"]
        finally:
            reset_current_tenant_id(token)

    with ThreadPoolExecutor(max_workers=2) as executor:
        ids = list(executor.map(lambda _: create_once(), range(2)))
    assert len(set(ids)) == 1


def test_a_workbook_is_read_by_the_worker_into_anchored_evidence(http: TestClient, tenant: str):
    import io

    from openpyxl import Workbook

    from backend.workers import source_worker

    project, process = _project(http)
    workbook = Workbook()
    workbook.active.title = "Ordini"
    workbook.active["B7"] = "Il CFO approva sopra 30.000"
    stream = io.BytesIO()
    workbook.save(stream)

    created = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={
            "roles": '["process_evidence"]',
            "retention": "persistent",
            "scopes": f'[{{"type":"process","id":"{process["id"]}"}}]',
        },
        files={"file": ("ordini.xlsx", stream.getvalue(), "application/octet-stream")},
    ).json()
    assert created["acquisition_status"] == "pending"
    assert http.get(f"/v1/workspace/sources/{created['id']}/evidence").json() == []

    _drain(source_worker, tenant)

    segments = http.get(f"/v1/workspace/sources/{created['id']}/evidence").json()
    assert [(item["ref"], item["text"]) for item in segments] == [("Ordini!B7", "Il CFO approva sopra 30.000")]
    assert segments[0]["locator"]["cell"] == "B7"
    document = http.get(f"/v1/workspace/sources/{created['id']}/document").json()
    assert "7: B=Il CFO approva sopra 30.000" in document["content"]


def test_an_unreadable_file_fails_once_a_service_outage_is_retried(http: TestClient, tenant: str, monkeypatch):
    from backend.workers import source_worker
    from backend.workspace_services.evidence import documents

    project, _ = _project(http)
    fields = {
        "roles": '["context"]',
        "retention": "persistent",
        "scopes": f'[{{"type":"project","id":"{project["id"]}"}}]',
    }
    import io

    from docx import Document

    word = Document()
    word.add_paragraph("Procedura")
    stream = io.BytesIO()
    word.save(stream)

    def outage(*_args, **_kwargs):
        raise documents.DocumentServiceUnavailable("Il servizio di lettura dei documenti non è disponibile.")

    monkeypatch.setattr(documents, "convert", outage)
    created = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data=fields,
        files={"file": (f"proc-{uuid.uuid4().hex[:6]}.docx", stream.getvalue(), "application/octet-stream")},
    ).json()
    _drain(source_worker, tenant)

    listed = {item["id"]: item for item in http.get(f"/v1/workspace/projects/{project['id']}/sources").json()}
    assert listed[created["id"]]["acquisition_status"] == "pending"
    assert "non è disponibile" in listed[created["id"]]["acquisition_error"]

    def broken(*_args, **_kwargs):
        raise documents.DocumentUnreadable("Il documento non può essere letto.")

    monkeypatch.setattr(documents, "convert", broken)
    from backend.workspace_database import workspace_connection
    from backend.workspace_storage import WorkspaceSource

    with workspace_connection() as session:
        session.get(WorkspaceSource, created["id"]).acquisition_next_attempt_at = "1970-01-01T00:00:00+00:00"
    _drain(source_worker, tenant)

    listed = {item["id"]: item for item in http.get(f"/v1/workspace/projects/{project['id']}/sources").json()}
    assert listed[created["id"]]["acquisition_status"] == "failed"

    # Ricaricare lo stesso file e' il "riprova": la fonte fallita torna in coda.
    monkeypatch.setattr(documents, "convert", lambda filename, payload: (_one_paragraph(), []))
    retried = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data=fields,
        files={"file": (created["name"], stream.getvalue(), "application/octet-stream")},
    )
    assert retried.status_code == 200
    assert retried.json()["id"] == created["id"]
    assert retried.json()["acquisition_status"] == "pending"
    _drain(source_worker, tenant)
    listed = {item["id"]: item for item in http.get(f"/v1/workspace/projects/{project['id']}/sources").json()}
    assert listed[created["id"]]["acquisition_status"] == "done"


def test_a_file_that_kills_the_worker_does_not_come_back_forever(http: TestClient, tenant: str):
    """Ogni presa in carico e' un tentativo: anche quella il cui worker muore
    prima di scrivere un esito, e che torna solo per scadenza del lease."""
    from backend import workspace_database as wd
    from backend.workspace_storage import WorkspaceSource

    project, _ = _project(http)
    created = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={
            "roles": '["context"]',
            "retention": "persistent",
            "scopes": f'[{{"type":"project","id":"{project["id"]}"}}]',
        },
        files={"file": (f"nota-{uuid.uuid4().hex[:6]}.txt", b"Una nota.", "text/plain")},
    ).json()

    for _ in range(wd.ACQUISITION_MAX_ATTEMPTS):
        claimed = wd.due_source_acquisitions(10, only_tenant_id=tenant)
        assert [row["id"] for row in claimed] == [created["id"]]
        # il worker muore qui: nessun esito, solo il lease che scade
        with wd.workspace_connection() as session:
            session.get(WorkspaceSource, created["id"]).acquisition_next_attempt_at = (
                "1970-01-01T00:00:00+00:00"
            )

    assert wd.due_source_acquisitions(10, only_tenant_id=tenant) == []
    listed = {item["id"]: item for item in http.get(f"/v1/workspace/projects/{project['id']}/sources").json()}
    assert listed[created["id"]]["acquisition_status"] == "failed"


def _one_paragraph():
    from docling_core.types.doc import DocItemLabel, DoclingDocument

    document = DoclingDocument(name="procedura")
    document.add_text(label=DocItemLabel.TEXT, text="Procedura\x00 con un byte NUL")
    return document


def test_a_chat_upload_removed_before_sending_is_discarded(http: TestClient, tenant: str):
    project, process = _project(http)
    fields = {
        "roles": '["process_evidence"]',
        "retention": "persistent",
        "scopes": f'[{{"type":"process","id":"{process["id"]}"}}]',
    }
    created = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data=fields,
        files={"file": (f"nota-{uuid.uuid4().hex[:6]}.txt", b"Da scartare", "text/plain")},
    ).json()

    assert http.delete(f"/v1/workspace/sources/{created['id']}").status_code == 204
    listed = [item["id"] for item in http.get(f"/v1/workspace/projects/{project['id']}/sources").json()]
    assert created["id"] not in listed
    assert http.get(f"/v1/workspace/sources/{created['id']}/original").status_code == 404
    assert http.delete(f"/v1/workspace/sources/{created['id']}").status_code == 404


def test_a_source_already_used_as_evidence_is_not_discarded(http: TestClient, tenant: str):
    from backend.workers import source_worker

    project, process = _project(http)
    created = http.post(
        f"/v1/workspace/projects/{project['id']}/sources/upload",
        data={
            "roles": '["process_evidence"]',
            "retention": "persistent",
            "scopes": f'[{{"type":"process","id":"{process["id"]}"}}]',
        },
        files={"file": (f"proc-{uuid.uuid4().hex[:6]}.txt", b"Il CFO approva", "text/plain")},
    ).json()
    _drain(source_worker, tenant)
    assert http.post(f"/v1/workspace/sources/{created['id']}/verify").status_code == 200

    refused = http.delete(f"/v1/workspace/sources/{created['id']}")
    assert refused.status_code == 409
