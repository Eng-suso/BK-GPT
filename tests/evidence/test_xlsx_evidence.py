"""Evidence Bucket per Excel: ogni cella arriva con la sua ancora, il suo tipo e la sua formula."""

from __future__ import annotations

import datetime as dt
import io
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table

from backend.workspace_services.evidence.xlsx import (
    WorkbookUnreadable,
    parse_xlsx,
    render_text,
    sheet_ref,
)
from backend.workspace_services.source_ingestion import SourceFileError, parse_source_file


def _save(workbook: Workbook) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _with_part(payload: bytes, name: str, content: str) -> bytes:
    """Aggiunge una parte all'archivio, come farebbe Excel con un disegno."""
    source = zipfile.ZipFile(io.BytesIO(payload))
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item))
        target.writestr(name, content)
    return stream.getvalue()


_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _hand_written_workbook(sheet_xml: str, *, date1904: bool = False) -> bytes:
    """Un workbook scritto come lo salva Excel, con i valori calcolati dentro.

    openpyxl scrive le formule ma non le calcola: per provare formule condivise
    e valori salvati serve un file in cui Excel abbia gia' fatto il suo lavoro.
    """
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>"
        ),
        "xl/workbook.xml": (
            f'<workbook xmlns="{_MAIN}" xmlns:r="{_REL}">'
            + ('<workbookPr date1904="1"/>' if date1904 else "")
            + '<sheets><sheet name="Ordini" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
            "</Relationships>"
        ),
        "xl/styles.xml": (
            f'<styleSheet xmlns="{_MAIN}">'
            '<fonts count="1"><font/></fonts><fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
            '<borders count="1"><border/></borders>'
            '<cellStyleXfs count="1"><xf/></cellStyleXfs>'
            '<cellXfs count="2"><xf numFmtId="0"/><xf numFmtId="14" applyNumberFormat="1"/></cellXfs>'
            "</styleSheet>"
        ),
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{_MAIN}"><sheetData>{sheet_xml}</sheetData></worksheet>',
    }
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return stream.getvalue()


def test_every_cell_keeps_its_own_column_when_the_row_has_gaps():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ordini"
    sheet["A1"] = "Ordine 42"
    sheet["C1"] = 1200
    sheet["E1"] = "CFO"

    source = parse_xlsx(_save(workbook))

    assert [item.anchor.ref for item in source.segments] == ["Ordini!A1", "Ordini!C1", "Ordini!E1"]
    assert source.segment("Ordini!C1").value == 1200
    assert source.segment("Ordini!C1").anchor.locator == {
        "sheet": "Ordini",
        "sheet_index": 0,
        "row": 1,
        "column": 3,
        "cell": "C1",
    }
    # Il vecchio estrattore saltava le celle vuote: "1200" finiva sotto la colonna B.
    assert "1: A=Ordine 42 | C=1200 | E=CFO" in render_text(source)


def test_shared_formulas_are_translated_and_carry_the_saved_value():
    source = parse_xlsx(
        _hand_written_workbook(
            '<row r="1"><c r="A1"><v>10</v></c><c r="B1"><f t="shared" ref="B1:B3" si="0">A1*2</f><v>20</v></c></row>'
            '<row r="2"><c r="A2"><v>11</v></c><c r="B2"><f t="shared" si="0"/><v>22</v></c></row>'
            '<row r="3"><c r="A3"><v>12</v></c><c r="B3"><f t="shared" si="0"/><v>24</v></c></row>'
            '<row r="4"><c r="B4"><f>SUM(B1:B3)/$A$1</f><v>6.6</v></c></row>'
        )
    )

    assert source.segment("Ordini!B1").attributes["formula"] == "A1*2"
    assert source.segment("Ordini!B3").attributes["formula"] == "A3*2"
    assert source.segment("Ordini!B3").value == 24
    assert source.segment("Ordini!B4").attributes["formula"] == "SUM(B1:B3)/$A$1"
    assert source.segment("Ordini!B4").value == 6.6
    assert source.status == "complete"
    assert "B=24 (=A3*2)" in render_text(source)


def test_a_formula_without_a_saved_value_is_declared_not_guessed():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = 5
    sheet["A2"] = "=A1*3"

    source = parse_xlsx(_save(workbook))
    segment = source.segment("Sheet!A2")

    assert segment.value_type == "empty"
    assert segment.attributes["formula"] == "A1*3"
    assert [issue.code for issue in source.issues] == ["formula_not_calculated"]
    assert source.issues[0].ref == "Sheet!A2"


def test_dates_come_out_as_dates_in_both_date_systems():
    in_1900 = parse_xlsx(_hand_written_workbook('<row r="1"><c r="A1" s="1"><v>46082</v></c></row>'))
    in_1904 = parse_xlsx(
        _hand_written_workbook('<row r="1"><c r="A1" s="1"><v>44620</v></c></row>', date1904=True)
    )

    assert in_1900.segment("Ordini!A1").value == "2026-03-01"
    assert in_1900.segment("Ordini!A1").value_type == "date"
    assert in_1900.structure["date_system"] == 1900
    assert in_1904.segment("Ordini!A1").value == "2026-03-01"
    assert in_1904.structure["date_system"] == 1904


def test_types_are_kept_and_a_typed_error_text_is_not_an_error():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = True
    sheet["A2"] = dt.datetime(2026, 3, 1, 9, 30)
    sheet["A3"] = 0.1

    source = parse_xlsx(_save(workbook))

    assert (source.segment("Sheet!A1").value_type, source.segment("Sheet!A1").value) == ("boolean", True)
    assert source.segment("Sheet!A2").value == "2026-03-01T09:30:00"
    assert source.segment("Sheet!A3").text == "0.1"

    # Excel salva "#N/A" scritto a mano come stringa, e il risultato di una
    # formula in errore come t="e": solo il secondo e' un errore.
    typed = parse_xlsx(
        _hand_written_workbook(
            '<row r="1"><c r="A1" t="inlineStr"><is><t>#N/A</t></is></c>'
            '<c r="B1" t="e"><f>1/0</f><v>#DIV/0!</v></c></row>'
        )
    )
    assert typed.segment("Ordini!A1").value_type == "text"
    assert typed.segment("Ordini!B1").value_type == "error"
    assert typed.segment("Ordini!B1").text == "#DIV/0!"


def test_merged_hidden_tables_and_comments_are_part_of_the_evidence():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ciclo passivo"
    sheet["A1"] = "Approvazioni 2026"
    sheet.merge_cells("A1:C1")
    sheet["A2"], sheet["B2"], sheet["C2"] = "Ordine", "Importo", "Approvatore"
    sheet["A3"], sheet["B3"], sheet["C3"] = "PO-1", 31000, "CFO"
    sheet["A4"], sheet["B4"], sheet["C4"] = "PO-2", 800, "Buyer"
    sheet.add_table(Table(displayName="Approvazioni", ref="A2:C4"))
    sheet["C3"].comment = Comment("Sopra 30k serve il CFO", "Rossi")
    sheet.row_dimensions[4].hidden = True
    archive = workbook.create_sheet("Storico")
    archive["A1"] = "vecchio"
    archive.sheet_state = "hidden"

    source = parse_xlsx(_save(workbook))

    title = source.segment("'Ciclo passivo'!A1")
    assert title.attributes["merged_range"] == "A1:C1"
    approver = source.segment("'Ciclo passivo'!C3")
    assert approver.attributes["table"] == "Approvazioni"
    assert approver.attributes["table_role"] == "data"
    assert approver.attributes["table_column"] == "Approvatore"
    assert source.segment("'Ciclo passivo'!A2").attributes["table_role"] == "header"
    assert source.segment("'Ciclo passivo'!B4").attributes["hidden"] is True
    assert source.segment("Storico!A1").attributes["sheet_hidden"] is True

    comment = source.segment("'Ciclo passivo'!C3#commento")
    assert comment.anchor.kind == "cell_comment"
    assert comment.text.startswith("Sopra 30k serve il CFO")
    assert comment.attributes["author"] == "Rossi"

    names = [sheet["name"] for sheet in source.structure["sheets"]]
    assert names == ["Ciclo passivo", "Storico"]
    assert "[Storico] (foglio nascosto)" in render_text(source)


def test_text_in_drawings_makes_the_source_partial():
    workbook = Workbook()
    workbook.active["A1"] = "Dati"
    payload = _with_part(
        _save(workbook),
        "xl/drawings/drawing1.xml",
        '<xdr:wsDr xmlns:xdr="urn:xdr" xmlns:a="urn:a"><a:t>Nota a margine</a:t></xdr:wsDr>',
    )

    source = parse_xlsx(payload)

    assert source.status == "partial"
    assert source.issues[0].code == "drawing_text_not_read"


def test_encrypted_and_unsafe_workbooks_are_refused_with_a_reason():
    with pytest.raises(WorkbookUnreadable, match="password"):
        parse_xlsx(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    workbook = Workbook()
    unsafe = _with_part(
        _save(workbook),
        "xl/custom.xml",
        '<!DOCTYPE x [<!ENTITY a "boom">]><x>&a;</x>',
    )
    with pytest.raises(WorkbookUnreadable, match="XML non sicuro"):
        parse_xlsx(unsafe)


def test_sheet_references_follow_excel_quoting():
    assert sheet_ref("Ordini", "B7") == "Ordini!B7"
    assert sheet_ref("Ciclo passivo", "B7") == "'Ciclo passivo'!B7"
    assert sheet_ref("Dell'Acqua", "B7") == "'Dell''Acqua'!B7"
    assert sheet_ref("AB12", "B7") == "'AB12'!B7"


def test_upload_parsing_carries_the_evidence_with_the_text():
    workbook = Workbook()
    workbook.active.title = "Ordini"
    workbook.active["B7"] = "Ordine 42"

    parsed = parse_source_file("registro.xlsx", _save(workbook), None)

    assert parsed.parser == "openpyxl-structural/1"
    assert parsed.evidence is not None
    assert parsed.evidence.segment("Ordini!B7").text == "Ordine 42"
    assert "7: B=Ordine 42" in parsed.text

    with pytest.raises(SourceFileError, match="non è valido"):
        parse_source_file("registro.xlsx", b"PK not really", None)
