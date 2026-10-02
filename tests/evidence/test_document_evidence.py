"""Evidence Bucket per PDF, DOCX e PPTX: docling-serve legge, docling-core porziona."""

from __future__ import annotations

import io

import httpx
import pypdfium2
import pytest
from docling_core.types.doc import (
    BoundingBox,
    ContentLayer,
    DocItemLabel,
    DoclingDocument,
    ProvenanceItem,
    Size,
    TableCell,
    TableData,
)

from backend.settings import settings
from backend.workspace_services.evidence import documents
from backend.workspace_services.source_ingestion import SourceFileError, parse_source_file


def _procedure() -> DoclingDocument:
    """Un documento come lo restituisce docling-serve per una procedura di due pagine."""
    document = DoclingDocument(name="procedura")
    document.add_page(page_no=1, size=Size(width=595, height=842))
    document.add_page(page_no=2, size=Size(width=595, height=842))

    def prov(page: int, top: float, length: int) -> ProvenanceItem:
        return ProvenanceItem(
            page_no=page,
            bbox=BoundingBox(l=72, t=top, r=400, b=top - 12),
            charspan=(0, length),
        )

    document.add_heading("Procedura acquisti", level=1, prov=prov(1, 780, 18))
    document.add_text(
        label=DocItemLabel.TEXT,
        text="Il CFO approva gli ordini sopra EUR 30.000.",
        prov=prov(1, 740, 43),
    )
    cells = [
        TableCell(text=text, start_row_offset_idx=row, end_row_offset_idx=row + 1,
                  start_col_offset_idx=col, end_col_offset_idx=col + 1, column_header=row == 0)
        for row, values in enumerate([["Soglia", "Approvatore"], ["30000", "CFO"]])
        for col, text in enumerate(values)
    ]
    document.add_table(data=TableData(num_rows=2, num_cols=2, table_cells=cells), prov=prov(2, 700, 0))
    document.add_text(
        label=DocItemLabel.TEXT,
        text="Verificare la soglia con Rossi",
        content_layer=ContentLayer.NOTES,
    )
    return document


def _service(document: DoclingDocument, *, task_status: str = "success", errors=None):
    """docling-serve finto, con le tre chiamate dell'API asincrona."""
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request.url.path)
        if request.url.path == "/v1/convert/file/async":
            return httpx.Response(200, json={"task_id": "t1", "task_status": "pending"})
        if request.url.path == "/v1/status/poll/t1":
            return httpx.Response(200, json={"task_id": "t1", "task_status": task_status})
        if request.url.path == "/v1/result/t1":
            return httpx.Response(
                200,
                json={
                    "status": "success",
                    "document": {"json_content": document.export_to_dict()},
                    "errors": errors or [],
                },
            )
        return httpx.Response(404)

    return httpx.MockTransport(handler), seen


def test_conversion_goes_through_the_async_api_and_returns_the_document():
    transport, seen = _service(_procedure())

    document, errors = documents.convert("procedura.pdf", b"%PDF-1.7", transport=transport)

    assert seen == ["/v1/convert/file/async", "/v1/status/poll/t1", "/v1/result/t1"]
    assert errors == []
    assert document.texts[1].text == "Il CFO approva gli ordini sopra EUR 30.000."


def test_every_chunk_cites_the_document_items_it_comes_from():
    source = documents.evidence_from(_procedure(), format="pdf")

    rule = next(item for item in source.segments if "30.000" in item.text)
    assert rule.anchor.ref == "#/texts/1"
    assert rule.anchor.locator["headings"] == ["Procedura acquisti"]
    assert rule.anchor.locator["provenance"][0]["page"] == 1

    table = next(item for item in source.segments if item.anchor.ref.startswith("#/tables/"))
    assert "CFO" in table.text
    assert table.anchor.locator["provenance"][0]["page"] == 2

    # Le note e i commenti sono evidenze: il chunker li include, marcati per livello.
    note = next(item for item in source.segments if "Rossi" in item.text)
    assert note.attributes["layers"] == ["notes"]
    assert source.status == "complete"


def test_service_errors_and_pictures_make_the_source_partial():
    document = _procedure()
    document.add_picture()

    source = documents.evidence_from(document, format="pdf", service_errors=["Pagina 2: tabella illeggibile"])

    assert source.status == "partial"
    assert [issue.code for issue in source.issues] == ["conversion_partial", "pictures_not_described"]


def test_an_unreachable_service_is_not_a_broken_file():
    def refuse(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("connection refused", request=request)

    with pytest.raises(documents.DocumentServiceUnavailable):
        documents.convert("procedura.pdf", b"%PDF-1.7", transport=httpx.MockTransport(refuse))

    transport, _ = _service(_procedure(), task_status="failure")
    with pytest.raises(documents.DocumentUnreadable):
        documents.convert("procedura.pdf", b"%PDF-1.7", transport=transport)


@pytest.mark.parametrize(
    "reply",
    [
        httpx.Response(200, json={"unexpected": True}),
        httpx.Response(200, text="<html>proxy error</html>"),
    ],
)
def test_an_unexpected_reply_is_a_service_problem_not_a_500(reply):
    transport = httpx.MockTransport(lambda request: reply)

    with pytest.raises(documents.DocumentServiceUnavailable):
        documents.convert("procedura.pdf", b"%PDF-1.7", transport=transport)


def test_a_file_refused_by_the_service_is_unreadable_not_an_outage():
    transport = httpx.MockTransport(lambda request: httpx.Response(415, json={"detail": "unsupported"}))
    with pytest.raises(documents.DocumentUnreadable):
        documents.convert("procedura.pdf", b"%PDF-1.7", transport=transport)

    busy = httpx.MockTransport(lambda request: httpx.Response(429, json={"detail": "slow down"}))
    with pytest.raises(documents.DocumentServiceUnavailable):
        documents.convert("procedura.pdf", b"%PDF-1.7", transport=busy)


def test_plain_string_errors_from_the_service_are_kept():
    transport, _ = _service(_procedure(), errors=["Pagina 3 illeggibile"])

    _document, errors = documents.convert("procedura.pdf", b"%PDF-1.7", transport=transport)

    assert errors == ["Pagina 3 illeggibile"]


def _pdf(pages: int) -> bytes:
    document = pypdfium2.PdfDocument.new()
    for _ in range(pages):
        document.new_page(595, 842)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_pdfs_are_checked_locally_before_reaching_the_service(monkeypatch):
    documents.check_pdf(_pdf(2))
    with pytest.raises(documents.DocumentUnreadable, match="PDF valido"):
        documents.check_pdf(b"not a pdf")
    with pytest.raises(documents.DocumentUnreadable, match="non può essere letto"):
        documents.check_pdf(b"%PDF-1.7 broken")
    monkeypatch.setattr(documents, "MAX_PDF_PAGES", 1)
    with pytest.raises(documents.DocumentUnreadable, match="limite"):
        documents.check_pdf(_pdf(2))


def test_upload_parsing_uses_the_document_service(monkeypatch):
    asked: dict = {}

    def fake_convert(filename, payload, **options):
        asked.update(options)
        return _procedure(), []

    monkeypatch.setattr(documents, "convert", fake_convert)

    parsed = parse_source_file("procedura.pdf", _pdf(1), "application/pdf")

    # Una pagina senza testo: e' una scansione, va letta con l'OCR.
    assert asked == {"do_ocr": True}

    assert parsed.parser == "docling-serve"
    assert "Il CFO approva gli ordini sopra EUR 30.000." in parsed.text
    assert parsed.evidence.segment("#/texts/1") is not None


def _service_is_up() -> bool:
    try:
        return httpx.get(settings.docling_serve_url.rstrip("/") + "/health", timeout=2).status_code == 200
    except httpx.HTTPError:
        return False


@pytest.mark.skipif(not _service_is_up(), reason="serve docling-serve (ops/docker-compose.yml)")
def test_real_service_reads_a_docx_with_headings_and_a_table():
    from docx import Document

    word = Document()
    word.add_heading("Procedura acquisti", 1)
    word.add_paragraph("Il CFO approva gli ordini sopra EUR 30.000.")
    table = word.add_table(rows=2, cols=2)
    for row, values in enumerate([["Soglia", "Approvatore"], ["30000", "CFO"]]):
        for col, text in enumerate(values):
            table.cell(row, col).text = text
    stream = io.BytesIO()
    word.save(stream)

    parsed = parse_source_file("procedura.docx", stream.getvalue(), None)

    rule = next(item for item in parsed.evidence.segments if "30.000" in item.text)
    assert rule.anchor.locator["headings"] == ["Procedura acquisti"]
    assert any("CFO" in item.text for item in parsed.evidence.segments if item.anchor.ref.startswith("#/tables/"))


@pytest.mark.skipif(not _service_is_up(), reason="serve docling-serve (ops/docker-compose.yml)")
def test_real_service_reads_slides_and_speaker_notes():
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Handoff"
    slide.placeholders[1].text = "Passaggio a Finance"
    slide.notes_slide.notes_text_frame.text = "Verificare con Rossi"
    stream = io.BytesIO()
    deck.save(stream)

    parsed = parse_source_file("handoff.pptx", stream.getvalue(), None)

    body = next(item for item in parsed.evidence.segments if "Passaggio a Finance" in item.text)
    assert body.anchor.locator["provenance"][0]["page"] == 1
    note = next(item for item in parsed.evidence.segments if "Rossi" in item.text)
    assert note.attributes["layers"] == ["notes"]


def test_unsafe_office_files_never_reach_the_service(monkeypatch):
    def fail(*_args, **_kwargs):
        raise AssertionError("il file non doveva arrivare al servizio")

    monkeypatch.setattr(documents, "convert", fail)
    import zipfile

    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY a "boom">]><w:document>&a;</w:document>')
    with pytest.raises(SourceFileError, match="XML non sicuro"):
        parse_source_file("procedura.docx", stream.getvalue(), None)


def _pdf_with_text(line: str) -> bytes:
    """Un PDF minimo con uno strato di testo vero, scritto a mano."""
    stream = f"BT /F1 12 Tf 72 720 Td ({line}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_ocr_is_asked_only_for_pdfs_without_a_text_layer():
    digital = _pdf_with_text("Il CFO approva gli ordini sopra EUR 30.000 dopo il controllo del budget.")
    assert documents.needs_ocr(digital) is False
    # Pagine senza testo: una scansione, si legge solo con l'OCR.
    assert documents.needs_ocr(_pdf(2)) is True


def test_the_service_is_told_whether_to_run_ocr():
    sent: list[str] = []
    transport, _seen = _service(_procedure())

    def recording(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/v1/convert/file/async":
            sent.append(request.content.decode("latin-1"))
        return transport.handle_request(request)

    documents.convert("procedura.pdf", b"%PDF-1.7", transport=httpx.MockTransport(recording), do_ocr=False)
    assert 'name="do_ocr"\r\n\r\nfalse' in sent[0]

